"""Memory-efficient XLSX Campaign Parser with streaming processing.

This service implements memory-efficient parsing of campaign XLSX files according to
campaignXLSXstructure.md specifications with the following features:

- Memory constraint: <100MB usage for 250MB files
- Streaming processing: 1000-record chunks for efficient memory usage
- File volume: Handle 7,057 campaign records per file
- Column count: Parse exactly 7 columns per specification
- Integration: Use CampaignsStaging model for database operations

Technical Requirements:
- German date format parsing (DD.MM.YYYY-DD.MM.YYYY)
- UUID validation and format compliance
- Null value handling for Budget field (8.7% rate)
- Comprehensive error handling with descriptive messages
- Integration with existing staging infrastructure
"""

import os
import gc
import psutil
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Optional, Iterator, Tuple, Callable
from uuid import UUID, uuid4
from pathlib import Path

# Third-party imports
import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Local imports
from src.models.staging import CampaignsStaging, CampaignStagingRequest, StagingValidationError
from src.database.staging_connection import StagingDatabase


class XLSXParsingError(Exception):
    """Exception raised during XLSX parsing operations with enhanced context support."""
    
    def __init__(self, message: str, row_number: Optional[int] = None, context: Optional[Dict[str, Any]] = None):
        super().__init__(message)
        self.message = message
        self.row_number = row_number
        self.context = context or {}
    
    def __str__(self):
        if self.row_number:
            return f"{self.message} in row {self.row_number}"
        return self.message


class XLSXCampaignParser:
    """Memory-efficient XLSX parser for campaign data with streaming processing."""
    
    # Expected column headers according to campaignXLSXstructure.md
    EXPECTED_HEADERS = [
        "Deal/Campaign name",
        "Runtime", 
        "Impression goal",
        "Budget €",
        "CPM €",
        "Deal/Campaign ID",
        "Buyer"
    ]
    
    # Default processing configuration
    DEFAULT_CHUNK_SIZE = 1000
    MAX_MEMORY_MB = 100
    
    def __init__(self, staging_db: Optional[StagingDatabase] = None):
        """Initialize the XLSX campaign parser.
        
        Args:
            staging_db: Optional staging database connection for saving data
        """
        self.staging_db = staging_db or StagingDatabase()
        self.chunk_size = self.DEFAULT_CHUNK_SIZE
        self.current_batch_id: Optional[UUID] = None
        self.processed_count = 0
        self.error_count = 0
    
    def get_default_chunk_size(self) -> int:
        """Get the default chunk size for processing."""
        return self.DEFAULT_CHUNK_SIZE
    
    def set_chunk_size(self, chunk_size: int) -> None:
        """Set the chunk size for processing."""
        if chunk_size <= 0:
            raise ValueError("Chunk size must be positive")
        self.chunk_size = chunk_size
    
    def get_chunk_size(self) -> int:
        """Get the current chunk size."""
        return self.chunk_size
    
    def generate_batch_id(self) -> UUID:
        """Generate a new batch ID for processing."""
        self.current_batch_id = uuid4()
        return self.current_batch_id
    
    def create_processing_metadata(self, batch_id: UUID, classification: str, row_number: int) -> Dict[str, Any]:
        """Create processing metadata for staging records."""
        return {
            'processing_batch_id': batch_id,
            'record_classification': classification,
            'source_row_number': row_number,
            'validation_errors': None
        }
    
    def parse_file_streaming(
        self, 
        file_path: str, 
        chunk_size: Optional[int] = None,
        skip_invalid: bool = False,
        progress_callback: Optional[Callable[[int, int], None]] = None
    ) -> Iterator[Dict[str, Any]]:
        """Parse XLSX file using streaming processing for memory efficiency.
        
        Args:
            file_path: Path to the XLSX file
            chunk_size: Optional chunk size override
            skip_invalid: Whether to skip invalid rows or raise errors
            progress_callback: Optional callback for progress reporting
            
        Yields:
            Dict containing parsed campaign data and processing metadata
            
        Raises:
            XLSXParsingError: For file format, validation, or processing errors
        """
        # Validate file exists
        if not os.path.exists(file_path):
            raise XLSXParsingError(f"File not found: {file_path}")
        
        # Set chunk size if provided
        if chunk_size:
            self.set_chunk_size(chunk_size)
        
        # Generate batch ID for this processing run
        batch_id = self.generate_batch_id()
        
        # Monitor memory usage
        process = psutil.Process(os.getpid())
        initial_memory = process.memory_info().rss / 1024 / 1024  # MB
        
        try:
            # Open workbook in read-only mode for memory efficiency
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            worksheet = workbook.active
            
            # Get all rows as iterator for memory efficiency
            row_iterator = worksheet.iter_rows(values_only=True)
            
            # Validate headers
            header_row = next(row_iterator, None)
            if not header_row:
                raise XLSXParsingError("Empty XLSX file")
            
            self._validate_headers(header_row)
            
            # Count total rows for progress reporting
            total_rows = worksheet.max_row - 1  # Exclude header
            self.processed_count = 0
            self.error_count = 0
            
            # Process data in chunks
            current_chunk = []
            
            for row_number, row_data in enumerate(row_iterator, start=2):  # Start from row 2 (after header)
                try:
                    # Skip empty rows
                    if not any(cell for cell in row_data if cell is not None):
                        continue
                    
                    # Parse individual row
                    parsed_record = self.parse_row(row_data, row_number)
                    
                    # Add processing metadata
                    metadata = self.create_processing_metadata(batch_id, "valid", row_number)
                    parsed_record.update(metadata)
                    
                    current_chunk.append(parsed_record)
                    self.processed_count += 1
                    
                    # Process chunk when full
                    if len(current_chunk) >= self.chunk_size:
                        yield from self._process_chunk(current_chunk)
                        current_chunk = []
                        
                        # Force garbage collection to maintain memory efficiency
                        gc.collect()
                        
                        # Check memory usage
                        current_memory = process.memory_info().rss / 1024 / 1024  # MB
                        memory_increase = current_memory - initial_memory
                        
                        if memory_increase > self.MAX_MEMORY_MB:
                            raise XLSXParsingError(
                                f"Memory usage exceeded limit: {memory_increase:.2f}MB > {self.MAX_MEMORY_MB}MB"
                            )
                    
                    # Report progress
                    if progress_callback:
                        progress_callback(self.processed_count, total_rows)
                
                except XLSXParsingError as e:
                    self.error_count += 1
                    
                    if skip_invalid:
                        # Log error and continue processing
                        error_record = {
                            'source_row_number': row_number,
                            'validation_errors': str(e),
                            'processing_batch_id': batch_id,
                            'record_classification': 'invalid'
                        }
                        yield error_record
                        continue
                    else:
                        raise
            
            # Process remaining chunk
            if current_chunk:
                yield from self._process_chunk(current_chunk)
            
            # Final progress report
            if progress_callback:
                progress_callback(self.processed_count, total_rows)
        
        except openpyxl.utils.exceptions.InvalidFileException:
            raise XLSXParsingError(f"Invalid XLSX file format: {file_path}")
        
        finally:
            # Ensure workbook is closed to free memory
            try:
                workbook.close()
            except:
                pass
    
    def _validate_headers(self, header_row: Tuple) -> None:
        """Validate that header row matches expected structure.
        
        Args:
            header_row: First row of the XLSX file
            
        Raises:
            XLSXParsingError: If headers don't match expected structure
        """
        if len(header_row) != len(self.EXPECTED_HEADERS):
            raise XLSXParsingError(
                f"Invalid XLSX structure: Expected {len(self.EXPECTED_HEADERS)} columns, got {len(header_row)}"
            )
        
        for i, (expected, actual) in enumerate(zip(self.EXPECTED_HEADERS, header_row)):
            if actual != expected:
                raise XLSXParsingError(
                    f"Invalid XLSX structure: Column {i+1} expected '{expected}', got '{actual}'"
                )
    
    def _process_chunk(self, chunk: List[Dict[str, Any]]) -> Iterator[Dict[str, Any]]:
        """Process a chunk of parsed records.
        
        Args:
            chunk: List of parsed campaign records
            
        Yields:
            Processed records with any additional metadata
        """
        # Save chunk to database if staging_db is available
        if self.staging_db:
            self._save_chunk_to_database(chunk)
        
        # Yield all records in chunk
        for record in chunk:
            yield record
    
    def _save_chunk_to_database(self, chunk: List[Dict[str, Any]]) -> None:
        """Save a chunk of records to the staging database.
        
        Args:
            chunk: List of campaign records to save
        """
        try:
            self.staging_db.save_campaigns_batch(chunk)
        except Exception as e:
            raise XLSXParsingError(f"Failed to save chunk to database: {str(e)}")
    
    def parse_row(self, row_data: Tuple, row_number: int) -> Dict[str, Any]:
        """Parse a single row of campaign data.
        
        Args:
            row_data: Tuple containing cell values for the row
            row_number: Row number for error reporting
            
        Returns:
            Dict containing parsed campaign data
            
        Raises:
            XLSXParsingError: For validation or parsing errors
        """
        try:
            if len(row_data) != len(self.EXPECTED_HEADERS):
                raise XLSXParsingError(
                    f"Row has {len(row_data)} columns, expected {len(self.EXPECTED_HEADERS)}",
                    row_number=row_number
                )
            
            # Parse each column according to specification
            deal_campaign_name = self.parse_campaign_name(row_data[0])
            runtime = row_data[1]
            start_date, end_date = self.parse_runtime(runtime)
            impression_goal = self.parse_impression_goal(row_data[2])
            budget_eur = self.parse_budget(row_data[3])
            cpm_eur = self.parse_cpm(row_data[4])
            deal_campaign_id = self.parse_deal_id(row_data[5])
            buyer = self.parse_buyer(row_data[6])
            
            return {
                'deal_campaign_name': deal_campaign_name,
                'runtime': runtime,
                'start_date': start_date,
                'end_date': end_date,
                'impression_goal': impression_goal,
                'budget_eur': budget_eur,
                'cpm_eur': cpm_eur,
                'deal_campaign_id': deal_campaign_id,
                'buyer': buyer
            }
        
        except Exception as e:
            # Re-raise with row context if not already an XLSXParsingError
            if isinstance(e, XLSXParsingError):
                if not e.row_number:
                    e.row_number = row_number
                raise
            else:
                raise XLSXParsingError(str(e), row_number=row_number)
    
    def parse_campaign_name(self, value: Any) -> str:
        """Parse Deal/Campaign name column (VARCHAR(500), 100% populated).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed campaign name
            
        Raises:
            XLSXParsingError: For validation errors
        """
        if value is None or str(value).strip() == "":
            raise XLSXParsingError("Deal/Campaign name cannot be empty")
        
        name = str(value).strip()
        
        if len(name) > 500:
            raise XLSXParsingError("Deal/Campaign name exceeds 500 characters")
        
        return name
    
    def parse_runtime(self, value: Any) -> Tuple[datetime, datetime]:
        """Parse Runtime column (German date format DD.MM.YYYY-DD.MM.YYYY).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Tuple of (start_date, end_date)
            
        Raises:
            XLSXParsingError: For parsing or validation errors
        """
        if value is None or str(value).strip() == "":
            raise XLSXParsingError("Runtime cannot be empty")
        
        runtime_str = str(value).strip()
        
        try:
            # Split on hyphen
            if '-' not in runtime_str:
                raise XLSXParsingError(f"Invalid runtime format: {runtime_str}. Expected 'DD.MM.YYYY-DD.MM.YYYY'")
            
            start_str, end_str = runtime_str.split('-', 1)
            start_str = start_str.strip()
            end_str = end_str.strip()
            
            # Parse German date format
            start_date = datetime.strptime(start_str, '%d.%m.%Y')
            end_date = datetime.strptime(end_str, '%d.%m.%Y')
            
            # Validate date logic
            if start_date >= end_date:
                raise XLSXParsingError(f"Invalid runtime format: {runtime_str}. Start date must be before end date")
            
            return start_date, end_date
        
        except ValueError as e:
            raise XLSXParsingError(f"Invalid runtime format: {runtime_str}. Expected 'DD.MM.YYYY-DD.MM.YYYY'") from e
    
    def parse_impression_goal(self, value: Any) -> int:
        """Parse Impression goal column (INTEGER, range 3,000-720,000).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed impression goal
            
        Raises:
            XLSXParsingError: For validation errors
        """
        if value is None:
            raise XLSXParsingError("Invalid impression goal: cannot be empty")
        
        try:
            # Handle string or numeric input
            if isinstance(value, str):
                # Remove any thousand separators or scientific notation
                value = value.replace(',', '').replace(' ', '')
                goal = int(float(value))  # Convert via float to handle scientific notation
            else:
                goal = int(value)
            
            # Validate range according to specification
            if goal < 3000 or goal > 720000:
                raise XLSXParsingError(f"Invalid impression goal: {goal}. Must be between 3,000 and 720,000")
            
            return goal
        
        except (ValueError, TypeError) as e:
            raise XLSXParsingError(f"Invalid impression goal: {value}. Must be a number between 3,000 and 720,000") from e
    
    def parse_budget(self, value: Any) -> Optional[Decimal]:
        """Parse Budget column (DECIMAL(12,2), 8.7% null rate, nullable).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed budget or None if null
            
        Raises:
            XLSXParsingError: For validation errors
        """
        # Handle null values (8.7% null rate)
        if value is None or str(value).strip() == "":
            return None
        
        try:
            budget = Decimal(str(value))
            
            # Validate range (must be non-negative)
            if budget < 0:
                raise XLSXParsingError(f"Invalid budget: {budget}. Must be non-negative")
            
            # Validate precision (max 12 digits total, 2 decimal places = max 10 integer digits)
            if budget >= Decimal('10000000000'):  # 10^10 (10 integer digits max for DECIMAL(12,2))
                raise XLSXParsingError(f"Invalid budget: {budget}. Exceeds maximum value")
            
            return budget
        
        except (InvalidOperation, ValueError) as e:
            raise XLSXParsingError(f"Invalid budget: {value}. Must be a valid decimal number") from e
    
    def parse_cpm(self, value: Any) -> Decimal:
        """Parse CPM column (DECIMAL(8,2), range €0.01-€45.00, 100% populated).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed CPM value
            
        Raises:
            XLSXParsingError: For validation errors
        """
        if value is None or str(value).strip() == "":
            raise XLSXParsingError("Invalid CPM: cannot be empty")
        
        try:
            cpm = Decimal(str(value))
            
            # Validate range according to specification
            if cpm < Decimal('0.01') or cpm > Decimal('45.00'):
                raise XLSXParsingError(f"Invalid CPM: {cpm}. Must be between €0.01 and €45.00")
            
            return cpm
        
        except (InvalidOperation, ValueError) as e:
            raise XLSXParsingError(f"Invalid CPM: {value}. Must be a valid decimal between €0.01 and €45.00") from e
    
    def parse_deal_id(self, value: Any) -> UUID:
        """Parse Deal/Campaign ID column (UUID format, 100% unique).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed UUID
            
        Raises:
            XLSXParsingError: For validation errors
        """
        if value is None or str(value).strip() == "":
            raise XLSXParsingError("Invalid Deal UUID format: cannot be empty")
        
        try:
            # Handle UUID or string input
            if isinstance(value, UUID):
                return value
            
            uuid_str = str(value).strip()
            return UUID(uuid_str)
        
        except (ValueError, TypeError) as e:
            raise XLSXParsingError(f"Invalid Deal UUID format: {value}. Must be a valid UUID") from e
    
    def parse_buyer(self, value: Any) -> str:
        """Parse Buyer column (VARCHAR(100), 100% populated).
        
        Args:
            value: Cell value to parse
            
        Returns:
            Parsed buyer name
            
        Raises:
            XLSXParsingError: For validation errors
        """
        if value is None or str(value).strip() == "":
            raise XLSXParsingError("Buyer cannot be empty")
        
        buyer = str(value).strip()
        
        if len(buyer) > 100:
            raise XLSXParsingError("Buyer name exceeds 100 characters")
        
        return buyer


# Factory function for easy instantiation
def create_xlsx_parser(staging_db: Optional[StagingDatabase] = None) -> XLSXCampaignParser:
    """Create an XLSX campaign parser instance.
    
    Args:
        staging_db: Optional staging database connection
        
    Returns:
        Configured XLSXCampaignParser instance
    """
    return XLSXCampaignParser(staging_db)


# Example usage and integration patterns
if __name__ == "__main__":
    # Example usage
    parser = create_xlsx_parser()
    
    try:
        # Process a campaign XLSX file
        for record in parser.parse_file_streaming("campaigns.xlsx", progress_callback=lambda p, t: print(f"Progress: {p}/{t}")):
            print(f"Processed campaign: {record['deal_campaign_name']}")
            
        print(f"Processing complete: {parser.processed_count} records processed, {parser.error_count} errors")
        
    except XLSXParsingError as e:
        print(f"Parsing error: {e}")
    except Exception as e:
        print(f"Unexpected error: {e}")